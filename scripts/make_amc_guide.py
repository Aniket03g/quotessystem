# -*- coding: utf-8 -*-
"""Builds the short Warranty & AMC guide as a .docx.

Same house style as make_lead_guide.py and make_room_pdf_guide.py — the guides
are read side by side, so the headings, callouts and tables are identical.
"""

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r'D:\GROVE SYSTEMS\working proxy - Leads\docs\Warranty and AMC Guide.docx'

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x5F, 0x6D)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

doc = docx.Document()

for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12

for name, size in (('Heading 1', 16), ('Heading 2', 12.5)):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = NAVY
    st.font.bold = True
    st.paragraph_format.space_before = Pt(15 if name == 'Heading 1' else 11)
    st.paragraph_format.space_after = Pt(5)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def no_borders(table_):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none')
        borders.append(el)
    table_._tbl.tblPr.append(borders)


def fixed_widths(table_, widths):
    """Word ignores column .width unless the table is fixed-layout and every
    cell carries the width itself."""
    table_.autofit = False
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    table_._tbl.tblPr.append(layout)
    for row in table_.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)


def para(text, size=11, bold=False, italic=False, color=BLACK, after=8, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def steps(items):
    """Numbered by hand: the List Number style keeps counting across the whole
    document, so a later list would start at 9 instead of 1."""
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.4)
        pf.first_line_indent = Inches(-0.25)
        pf.space_after = Pt(5)
        n = p.add_run('%d.  ' % i)
        n.bold = True
        n.font.size = Pt(11)
        for text, bold in (it if isinstance(it, list) else [(it, False)]):
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(11)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        for text, bold in (it if isinstance(it, list) else [(it, False)]):
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(11)
        p.paragraph_format.space_after = Pt(3)


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, '1E3A5F')
        p = c.paragraphs[0]
        p.text = ''
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2 == 1:
                shade(cells[i], 'F4F6F9')
            p = cells[i].paragraphs[0]
            p.text = ''
            r = p.add_run(val)
            r.font.size = Pt(10)
            r.bold = (i == 0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
    fixed_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(title, body, fill='E8F1FB', bar='2E6DA4'):
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    fixed_widths(t, [0.09, 6.31])
    shade(t.cell(0, 0), bar)
    shade(t.cell(0, 1), fill)
    c = t.cell(0, 1)
    p = c.paragraphs[0]
    p.text = ''
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)
    pb = c.add_paragraph()
    rb = pb.add_run(body)
    rb.font.size = Pt(10.5)
    pb.paragraph_format.space_after = Pt(4)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ===================================================================== TITLE ==
para('Grove Systems CRM', size=10, bold=True, color=GREY, after=2)
para('Warranty & AMC', size=24, bold=True, color=NAVY, after=2)
para('Knowing what is about to run out of cover', size=11.5, color=GREY, after=14)

para(
    'The Warranty & AMC page is the register of equipment we still cover \u2014 who '
    'owns it, where it is installed, which serial numbers it covers and the date '
    'the cover ends. Its job is to give us a couple of months\u2019 warning before '
    'cover lapses, so a renewal or an AMC can be quoted while the customer is '
    'still covered rather than after.',
    after=12,
)

callout(
    'Who can see this page',
    'Warranty & AMC is visible to admins and super admins. Both see the whole '
    'register, not just their own accounts. Sales users do not see the page or '
    'the sidebar link.',
)

# ================================================================== THE PAGE ==
doc.add_heading('1. Finding your way around', level=1)

para(
    'Open ' 'Warranty & AMC' ' from the sidebar. Records are listed soonest to '
    'expire first, so the ones that need attention are always at the top.',
    after=8,
)

para('The four boxes across the top are counts, and each one is a filter \u2014 click it to show only those records.', after=8)

table(
    ['Box', 'What it counts'],
    [
        ['Lapsed', 'Cover has already ended. These need chasing first.'],
        ['Expiring soon', 'Cover ends within the next 60 days.'],
        ['Active', 'More than 60 days of cover left. Nothing to do yet.'],
        ['Assets covered', 'How many individual units (serial numbers) the whole register covers.'],
    ],
    [1.6, 4.8],
)

para('The list itself shows one row per contract:', after=6)

table(
    ['Column', 'What it shows'],
    [
        ['Customer', 'The account the equipment belongs to.'],
        ['Location', 'Where it is installed.'],
        ['Expires', 'The date cover ends.'],
        ['Status', 'A coloured pill \u2014 "Lapsed 34d ago" or "58d left". Red and amber rows need action.'],
        ['Assets', 'How many units this record covers.'],
        ['Value', 'The value of the original contract.'],
        ['Expiry source', 'Where the expiry date came from. See section 3 \u2014 this one matters.'],
    ],
    [1.6, 4.8],
)

para(
    'Click any row to expand it in place. The serial numbers covered by that '
    'record appear as chips, along with the invoice details. Those serials are '
    'what an AMC quote is built from, so this is where you copy them from.',
    after=10,
)

# ================================================================== FILTERING ==
doc.add_heading('2. Finding a particular record', level=1)

bullets([
    [('The ', False), ('search box', True),
     (' matches customer, location, invoice number, model and serial number. '
      'If a customer calls about one unit, paste its serial in here.', False)],
    [('All cover', True), (' filters by Lapsed / Expiring in 60 days / Active \u2014 the same bands as the boxes above.', False)],
    [('Any expiry source', True), (' filters by AMS registered or Computed.', False)],
    [('Sort', True), (' by expiry, by contract value, or by customer name.', False)],
    [('Reset', True), (' clears everything and puts the list back to soonest-expiry first.', False)],
])

# ===================================================================== SOURCE ==
doc.add_heading('3. "AMS registered" vs "Computed"', level=1)

para(
    'Every record carries a badge saying where its expiry date came from. The '
    'short version: AMS registered is the real date, Computed is our safe '
    'estimate, and a Computed date is always early, never late.',
    after=8,
)

table(
    ['Badge', 'What it means'],
    [
        ['AMS registered',
         'The expiry was read from the Yealink AMS portal, where the unit\u2019s warranty is '
         'actually recorded. This is the exact date. Trust it.'],
        ['Computed',
         'The unit was never looked up in the AMS portal, so the CRM calculates the expiry '
         'as invoice date + 2 years. Warranty registration always happens some time after '
         'invoicing, so the true expiry is later than the date shown \u2014 typically by three '
         'to six months.'],
    ],
    [1.6, 4.8],
)

callout(
    'Why we would rather be early',
    'A Computed record will prompt you a few months before cover actually ends. '
    'That is deliberate. Chasing a renewal early costs a phone call; finding out '
    'a week after cover lapsed costs the renewal. If you need the exact date on a '
    'Computed record, look the serial up in the AMS portal \u2014 the record can then '
    'be switched to AMS registered.',
)

para(
    'Most of the register is Computed today, so treat "early" as the normal case '
    'rather than an exception. Use the Expiry source filter if you want to see '
    'only the dates that are known to be exact.',
    after=6,
)

callout(
    'Do not go back to the old spreadsheet',
    'The expiry column in the original AMC spreadsheet was typed by hand and '
    'contains mistakes \u2014 on at least one row the day and month are swapped. The '
    'CRM ignores that column completely and works from the invoice date and the '
    'AMS portal instead. The page is the source of truth, not the sheet.',
    fill='FDF3E7', bar='C97B2E',
)

# =================================================================== REMINDER ==
doc.add_heading('4. The reminder email', level=1)

bullets([
    'One email goes out 60 days before cover ends, listing the records that are approaching expiry.',
    'It is sent to every admin and super admin, not to an individual account owner.',
    [('Each record is reminded ', False), ('once', True),
     (' \u2014 the CRM records that it has been sent, so the same contract does not '
      'nag every day.', False)],
    'Anything already lapsed is picked up by the next scan rather than skipped, so nothing falls through.',
])

para(
    'The email is a prompt, not the work. When one arrives, open the page, expand '
    'the record and quote the renewal.',
    after=10,
)

# ================================================================== WHAT NEXT ==
doc.add_heading('5. What to do with an expiring record', level=1)

steps([
    'Open Warranty & AMC and click the record so the serial numbers are visible.',
    [('Raise a quote for the customer the usual way. An AMC quote is priced ', False),
     ('per serial number', True),
     (' \u2014 one line per unit, described as the cover being sold, for example '
      '"1 Year Extended Warranty for UVC86".', False)],
    'Put the serial number and the cover start and end dates on each line, so the customer can see exactly what is being renewed.',
    'Send the quote as normal. When it is accepted, the renewed cover has to be '
    'added to the register so its own expiry is tracked — tell the admin team so '
    'the new record is created.',
])

callout(
    'Remember what Computed means when you quote',
    'On a Computed record the cover probably runs a few months longer than the '
    'date shown. Confirm the real end date with the customer or the AMS portal '
    'before printing dates on a quote.',
)

doc.save(OUT)
print('written:', OUT)
